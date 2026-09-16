from datetime import datetime
from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from .. import models
from ..database import get_db
from ..auth import get_current_user
from ..membership import YEAR_RE, membership_events as _membership_events

router = APIRouter(prefix="/api/reports", tags=["reports"])


@router.get("/membership-years")
def membership_years(
    db: Session = Depends(get_db),
    current_user: str = Depends(get_current_user),
):
    """Distinct years for which a '<year> Membership' event exists (newest first)."""
    events = (
        db.query(models.Event)
        .filter(models.Event.eventName.ilike("%membership%"))
        .all()
    )
    years = set()
    for e in events:
        m = YEAR_RE.search(e.eventName or "")
        if m:
            years.add(int(m.group(0)))
    return {"years": sorted(years, reverse=True)}


@router.get("/unpaid-membership")
def unpaid_membership(
    year: int = Query(..., ge=1900, le=2100),
    db: Session = Depends(get_db),
    current_user: str = Depends(get_current_user),
):
    """Active, non-life members with no contribution to the given year's
    Membership event. Read-only — does not modify any data."""
    events = _membership_events(db, year)
    if not events:
        raise HTTPException(
            status_code=404,
            detail=f"No membership event found for {year} "
                   f"(expected an event like '{year} Membership').",
        )
    event_ids = [e.eventId for e in events]

    paid_subq = (
        db.query(models.Contribution.personId)
        .filter(models.Contribution.eventId.in_(event_ids))
        .distinct()
    )

    members = (
        db.query(models.Member)
        .filter(
            models.Member.status == "Active",
            (models.Member.lifeMember.is_(False)) | (models.Member.lifeMember.is_(None)),
            ~models.Member.personId.in_(paid_subq),
        )
        .order_by(models.Member.lastName, models.Member.firstName)
        .all()
    )

    return {
        "year": year,
        "events": [{"eventId": e.eventId, "eventName": e.eventName} for e in events],
        "count": len(members),
        "members": [
            {
                "personId": m.personId,
                "firstName": m.firstName,
                "lastName": m.lastName,
                "spouse": m.spouse,
                "email": m.email,
                "cellPhone": m.cellPhone,
                "homePhone": m.homePhone,
                "address1": m.address1,
                "address2": m.address2,
                "city": m.city,
                "state": m.state,
                "zip": m.zip,
            }
            for m in members
        ],
    }


# ── Member Directory ──────────────────────────────────────────────────────────

SCOPES = {
    "all":    "All Members",
    "active": "Active Members",
    "life":   "Life Members",
}

# (attribute on models.Member, column heading, excel column width)
DIRECTORY_COLUMNS = [
    ("lastName",   "Last Name",   18),
    ("firstName",  "First Name",  18),
    ("middleName", "Middle Name", 14),
    ("spouse",     "Spouse",      20),
    ("children",   "Children",    30),
    ("address1",   "Address 1",   28),
    ("address2",   "Address 2",   18),
    ("city",       "City",        18),
    ("state",      "State",        7),
    ("zip",        "Zip",         10),
    ("homePhone",  "Home Phone",  16),
    ("cellPhone",  "Cell Phone",  16),
    ("cellPhone2", "Cell Phone 2", 16),
    ("email",      "Email",       32),
    ("status",     "Status",      10),
    ("lifeMember", "Life Member", 12),
]


def _directory_members(db: Session, scope: str):
    """Members for the requested scope, sorted by last then first name."""
    if scope not in SCOPES:
        raise HTTPException(
            status_code=400,
            detail=f"Invalid scope '{scope}'. Expected one of: {', '.join(SCOPES)}.",
        )

    q = db.query(models.Member)
    if scope == "active":
        q = q.filter(models.Member.status == "Active")
    elif scope == "life":
        q = q.filter(models.Member.lifeMember.is_(True))

    return q.order_by(models.Member.lastName, models.Member.firstName).all()


def _member_row(m: models.Member) -> dict:
    row = {attr: getattr(m, attr) for attr, _, _ in DIRECTORY_COLUMNS}
    row["personId"] = m.personId
    return row


@router.get("/member-directory")
def member_directory(
    scope: str = Query("all", pattern="^(all|active|life)$"),
    db: Session = Depends(get_db),
    current_user: str = Depends(get_current_user),
):
    """Directory of members — spouse, children, address, phones and email.
    Read-only — does not modify any data."""
    members = _directory_members(db, scope)
    return {
        "scope": scope,
        "scopeLabel": SCOPES[scope],
        "count": len(members),
        "columns": [{"key": attr, "label": label} for attr, label, _ in DIRECTORY_COLUMNS],
        "members": [_member_row(m) for m in members],
    }


@router.get("/member-directory.xlsx")
def member_directory_xlsx(
    scope: str = Query("all", pattern="^(all|active|life)$"),
    db: Session = Depends(get_db),
    current_user: str = Depends(get_current_user),
):
    """Same directory as /member-directory, returned as an .xlsx workbook."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    members = _directory_members(db, scope)

    wb = Workbook()
    ws = wb.active
    ws.title = "Member Directory"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E79")

    ws.append([label for _, label, _ in DIRECTORY_COLUMNS])
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    def cell_value(m, attr):
        value = getattr(m, attr)
        if attr == "lifeMember":
            return "Yes" if value else "No"
        return "" if value is None else value

    for m in members:
        ws.append([cell_value(m, attr) for attr, _, _ in DIRECTORY_COLUMNS])

    for idx, (_, _, width) in enumerate(DIRECTORY_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(DIRECTORY_COLUMNS))}{ws.max_row}"

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"member_directory_{scope}_{datetime.now():%Y%m%d}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
